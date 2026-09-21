from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.core.paginator import Paginator
from products.models import (
    Machine,
    MachineImage,
    Category,
    Contact,
)

from .forms import MachineForm, MachineImageForm
from django.http import HttpResponse
import openpyxl
from openpyxl.styles import Font

# ===========================
# LOGIN
# ===========================

def dashboard_login(request):

    if request.user.is_authenticated:
        return redirect("dashboard:home")

    if request.method == "POST":

        username = request.POST.get("username")
        password = request.POST.get("password")

        user = authenticate(
            request,
            username=username,
            password=password
        )

        if user is not None and user.is_staff:

            login(request, user)

            messages.success(
                request,
                f"Welcome {user.username}"
            )

            return redirect("dashboard:home")

        messages.error(
            request,
            "Invalid Username or Password."
        )

    return render(
        request,
        "dashboard/login.html"
    )


# ===========================
# LOGOUT
# ===========================

@login_required
def dashboard_logout(request):

    logout(request)

    messages.success(
        request,
        "Logged out successfully."
    )

    return redirect("dashboard:login")


# ===========================
# DASHBOARD
# ===========================

@login_required
def dashboard_home(request):

    context = {

        "total_products": Machine.objects.count(),

        "total_categories": Category.objects.count(),

        "total_gallery": MachineImage.objects.count(),

        "total_contacts": Contact.objects.count(),

        "recent_products": Machine.objects.order_by(
            "-created_at"
        )[:5],

        "recent_contacts": Contact.objects.order_by(
            "-created_at"
        )[:5],

    }

    return render(
        request,
        "dashboard/dashboard.html",
        context
    )


# ===========================
# PRODUCT LIST
# ===========================

@login_required
def product_list(request):

    products = Machine.objects.all()

    search = request.GET.get("search")

    if search:
        products = products.filter(
            name__icontains=search
        )

    context = {

        "products": products.order_by("-created_at"),

        "search": search,

    }

    return render(
        request,
        "dashboard/product_list.html",
        context
    )


# ===========================
# ADD PRODUCT
# ===========================

@login_required
def product_create(request):

    if request.method == "POST":

        form = MachineForm(
            request.POST,
            request.FILES
        )

        if form.is_valid():

            form.save()

            messages.success(
                request,
                "Machine added successfully."
            )

            return redirect("dashboard:product_list")

    else:

        form = MachineForm()

    return render(
        request,
        "dashboard/product_form.html",
        {
            "form": form,
            "title": "Add Machine"
        }
    )


# ===========================
# UPDATE PRODUCT
# ===========================

@login_required
def product_update(request, pk):

    product = get_object_or_404(
        Machine,
        pk=pk
    )

    if request.method == "POST":

        form = MachineForm(
            request.POST,
            request.FILES,
            instance=product
        )

        if form.is_valid():

            form.save()

            messages.success(
                request,
                "Machine updated successfully."
            )

            return redirect(
                "dashboard:product_list"
            )

    else:

        form = MachineForm(
            instance=product
        )

    return render(
        request,
        "dashboard/product_form.html",
        {
            "form": form,
            "product": product,
            "title": "Edit Machine"
        }
    )


# ===========================
# DELETE PRODUCT
# ===========================

@login_required
def product_delete(request, pk):

    product = get_object_or_404(
        Machine,
        pk=pk
    )

    if request.method == "POST":

        product.delete()

        messages.success(
            request,
            "Machine deleted."
        )

        return redirect(
            "dashboard:product_list"
        )

    return render(
        request,
        "dashboard/product_delete.html",
        {
            "product": product
        }
    )


# ===========================
# GALLERY
# ===========================

@login_required
def product_gallery(request, pk):

    machine = get_object_or_404(
        Machine,
        pk=pk
    )

    images = machine.gallery.all()

    if request.method == "POST":

        form = MachineImageForm(
            request.POST,
            request.FILES
        )

        if form.is_valid():

            gallery = form.save(
                commit=False
            )

            gallery.machine = machine

            gallery.save()

            messages.success(
                request,
                "Image uploaded successfully."
            )

            return redirect(
                "dashboard:product_gallery",
                pk=pk
            )

    else:

        form = MachineImageForm()

    return render(
        request,
        "dashboard/gallery.html",
        {
            "machine": machine,
            "images": images,
            "form": form,
        }
    )


# ===========================
# CONTACTS
# ===========================




@login_required
def contact_list(request):

    contacts = Contact.objects.all().order_by("-created_at")

    # -------------------------------
    # Search
    # -------------------------------

    search = request.GET.get("search")

    if search:

        contacts = contacts.filter(

            full_name__icontains=search

        ) | Contact.objects.filter(

            email__icontains=search

        ) | Contact.objects.filter(

            phone__icontains=search

        )

    # -------------------------------
    # Filter
    # -------------------------------

    status = request.GET.get("status")

    if status == "read":

        contacts = contacts.filter(is_read=True)

    elif status == "unread":

        contacts = contacts.filter(is_read=False)

    # -------------------------------
    # Pagination
    # -------------------------------

    paginator = Paginator(contacts, 10)

    page_number = request.GET.get("page")

    page_obj = paginator.get_page(page_number)

    context = {

        "contacts": page_obj,

        "page_obj": page_obj,

        "search": search,

        "status": status,

    }

    return render(

        request,

        "dashboard/contact_list.html",

        context

    )


# ===========================
# PROFILE
# ===========================

@login_required
def profile(request):

    return render(
        request,
        "dashboard/profile.html"
    )
@login_required
def mark_contact_read(request, pk):

    contact = get_object_or_404(

        Contact,

        pk=pk

    )

    contact.is_read = True

    contact.save()

    messages.success(

        request,

        "Marked as Read."

    )

    return redirect("dashboard:contact_list")


@login_required
def mark_contact_unread(request, pk):

    contact = get_object_or_404(

        Contact,

        pk=pk

    )

    contact.is_read = False

    contact.save()

    messages.success(

        request,

        "Marked as Unread."

    )

    return redirect("dashboard:contact_list")


@login_required
def contact_delete(request, pk):

    contact = get_object_or_404(

        Contact,

        pk=pk

    )

    contact.delete()

    messages.success(

        request,

        "Enquiry deleted successfully."

    )

    return redirect("dashboard:contact_list")
@login_required
def export_contacts_excel(request):

    workbook = openpyxl.Workbook()

    sheet = workbook.active

    sheet.title = "Contact Enquiries"

    headers = [

        "Name",

        "Company",

        "Email",

        "Phone",

        "Machine",

        "Status",

        "Date",

        "Message",

    ]

    for col_num, header in enumerate(headers, 1):

        cell = sheet.cell(row=1, column=col_num)

        cell.value = header

        cell.font = Font(bold=True)

    contacts = Contact.objects.all().order_by("-created_at")

    row = 2

    for contact in contacts:

        sheet.cell(row=row, column=1).value = contact.full_name

        sheet.cell(row=row, column=2).value = contact.company_name

        sheet.cell(row=row, column=3).value = contact.email

        sheet.cell(row=row, column=4).value = contact.phone

        sheet.cell(
            row=row,
            column=5
        ).value = str(contact.machine) if contact.machine else ""

        sheet.cell(
            row=row,
            column=6
        ).value = "Read" if contact.is_read else "Unread"

        sheet.cell(
            row=row,
            column=7
        ).value = contact.created_at.strftime("%d-%m-%Y %H:%M")

        sheet.cell(
            row=row,
            column=8
        ).value = contact.message

        row += 1

    response = HttpResponse(

        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    )

    response["Content-Disposition"] = (

        'attachment; filename="contact_enquiries.xlsx"'

    )

    workbook.save(response)

    return response

# ===========================
# CATEGORY MANAGEMENT
# ===========================

@login_required
def category_list(request):

    # ADD CATEGORY
    if request.method == "POST":

        name = request.POST.get("name", "").strip()

        # Add
        if "add_category" in request.POST:

            if not name:
                messages.error(
                    request,
                    "Category name is required."
                )

            elif Category.objects.filter(
                name__iexact=name
            ).exists():

                messages.error(
                    request,
                    "This category already exists."
                )

            else:

                Category.objects.create(
                    name=name
                )

                messages.success(
                    request,
                    "Category added successfully."
                )

        # EDIT CATEGORY
        elif "edit_category" in request.POST:

            category_id = request.POST.get("category_id")

            category = get_object_or_404(
                Category,
                pk=category_id
            )

            if not name:

                messages.error(
                    request,
                    "Category name is required."
                )

            elif Category.objects.filter(
                name__iexact=name
            ).exclude(
                pk=category_id
            ).exists():

                messages.error(
                    request,
                    "This category already exists."
                )

            else:

                category.name = name
                category.save()

                messages.success(
                    request,
                    "Category updated successfully."
                )

        # DELETE CATEGORY
        elif "delete_category" in request.POST:

            category_id = request.POST.get("category_id")

            category = get_object_or_404(
                Category,
                pk=category_id
            )

            category.delete()

            messages.success(
                request,
                "Category deleted successfully."
            )

        return redirect("dashboard:category_list")

    categories = Category.objects.all().order_by("name")

    return render(
        request,
        "dashboard/category_list.html",
        {
            "categories": categories
        }
    )